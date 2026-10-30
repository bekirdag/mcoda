"""Offline, bounded document operations. Input/output paths are fixed by the host."""
import hashlib
import importlib.metadata
import io
import json
import math
import resource
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

def lower_limit(kind, maximum):
    inherited = resource.getrlimit(kind)
    bound = min([maximum] + [value for value in inherited if value != resource.RLIM_INFINITY])
    resource.setrlimit(kind, (bound, bound))
    return bound


# Darwin exposes RLIMIT_AS but does not provide Linux's address-space limit.
# Keep the supported CPU/file caps and disclose the missing memory cap.
LIMITS = {
    "addressSpaceBytes": lower_limit(resource.RLIMIT_AS, 768 * 1024 * 1024) if sys.platform == "linux" else None,
    "cpuSeconds": lower_limit(resource.RLIMIT_CPU, 10),
    "fileBytes": lower_limit(resource.RLIMIT_FSIZE, 16 * 1024 * 1024),
}
VERSIONS = {name: importlib.metadata.version(name) for name in ["python-docx", "openpyxl", "pypdf", "reportlab", "defusedxml", "charset-normalizer", "et_xmlfile", "lxml", "pillow", "typing_extensions"]}
MIME = {"md": "text/markdown", "pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def bounded_zip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
            raise ValueError("Document archive expands beyond the extraction limit")
        if any(entry.file_size > 8 * 1024 * 1024 or entry.flag_bits & 1 for entry in entries):
            raise ValueError("Oversized or encrypted archive member")
        names = [entry.filename for entry in entries]
        if len(set(names)) != len(names) or any(name.startswith("/") or ".." in name.split("/") for name in names):
            raise ValueError("Unsafe or ambiguous document archive")


def scalar(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def read_document(data, fmt, args):
    chars = args.get("maxChars", 16000)
    max_pages = args.get("maxPages", 10)
    max_rows = args.get("maxRows", 100)
    max_sheets = args.get("maxSheets", 4)
    result = {"format": fmt, "truncated": False, "warnings": []}
    remaining = chars

    def text(value):
        nonlocal remaining
        value = str(value)
        result["truncated"] |= len(value) > remaining
        selected = value[:remaining]
        remaining -= len(selected)
        return selected

    if fmt == "md":
        result["text"] = text(data.decode("utf-8", errors="strict"))
    elif fmt == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires an independently decrypted local copy")
        result["pageCount"] = len(reader.pages)
        result["pages"] = []
        result["truncated"] |= len(reader.pages) > max_pages
        for number, page in enumerate(reader.pages[:max_pages], 1):
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > 8 * 1024 * 1024:
                raise ValueError("PDF page content exceeds the extraction limit")
            value = page.extract_text() or ""
            result["pages"].append({"page": number, "text": text(value), "status": "text_extracted" if value.strip() else "no_extractable_text"})
            if not value.strip():
                result["warnings"].append(f"Page {number} has no extractable text; OCR or visual inspection may be needed. No OCR was performed.")
        result["warnings"].append("PDF text order and table layout are heuristic; verify the source page for ambiguous content.")
    elif fmt == "docx":
        from docx import Document
        from docx.table import Table
        bounded_zip(data)
        document = Document(io.BytesIO(data))
        result["blocks"] = []
        for index, block in enumerate(document.iter_inner_content()):
            if index >= 200 or remaining <= 0:
                result["truncated"] = True
                break
            if isinstance(block, Table):
                rows = [[text(cell.text) for cell in row.cells[:32]] for row in block.rows[:max_rows]]
                result["truncated"] |= len(block.rows) > max_rows or any(len(row.cells) > 32 for row in block.rows[:max_rows])
                result["blocks"].append({"type": "table", "rows": rows})
            else:
                result["blocks"].append({"type": "paragraph", "text": text(block.text)})
        result["warnings"].append("Main-body paragraphs/tables only; images, headers, footnotes and tracked changes are not a complete visual interpretation.")
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        bounded_zip(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            result["sheetCount"] = len(workbook.worksheets)
            result["sheets"] = []
            result["truncated"] |= len(workbook.worksheets) > max_sheets
            for sheet in workbook.worksheets[:max_sheets]:
                rows, formulas = [], []
                for row in sheet.iter_rows(max_row=min(sheet.max_row or max_rows, max_rows), max_col=min(sheet.max_column or 32, 32)):
                    if remaining <= 0:
                        result["truncated"] = True
                        break
                    values = []
                    for cell in row:
                        value = scalar(cell.value)
                        if isinstance(value, str):
                            value = text(value)
                        if cell.data_type == "f":
                            formulas.append({"cell": cell.coordinate, "expression": value})
                        values.append(value)
                    rows.append(values)
                result["truncated"] |= (sheet.max_row or 0) > max_rows or (sheet.max_column or 0) > 32
                result["sheets"].append({"name": sheet.title, "rows": rows, "formulas": formulas, "formulaValuesVerified": False})
            result["warnings"].append("Formula expressions are source text, never evaluated values. Macros and external links are not executed.")
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported document format")
    return result


def create_document(fmt, args):
    blocks = args.get("blocks", [])
    sheets = args.get("sheets", [])
    title = args.get("title", "")
    if fmt == "xlsx":
        if blocks or not sheets:
            raise ValueError("XLSX requires sheets and cannot consume text blocks")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        workbook = Workbook()
        workbook.remove(workbook.active)
        seen = set()
        for data in sheets:
            name = data["name"]
            if name.lower() in seen:
                raise ValueError("Worksheet names must be unique ignoring case")
            seen.add(name.lower())
            sheet = workbook.create_sheet(name)
            for row in data["rows"]:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # Literal input never becomes a formula or DDE expression.
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            if data.get("header", True) and data["rows"]:
                for cell in sheet[1]:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="234364")
                sheet.freeze_panes = "A2"
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
    if sheets or not blocks:
        raise ValueError("Text documents require blocks and cannot consume sheets")
    if fmt == "md":
        lines = ["# " + title, ""] if title else []
        safe_cell = lambda value: str(value).replace("|", "\\|").replace("\n", "<br>")
        for block in blocks:
            kind = block["type"]
            if kind == "table":
                headers = block["headers"]
                lines += ["| " + " | ".join(map(safe_cell, headers)) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
                lines += ["| " + " | ".join(map(safe_cell, row)) + " |" for row in block["rows"]]
            else:
                prefix = "#" * block.get("level", 2) + " " if kind == "heading" else "- " if kind == "bullet" else ""
                lines.append(prefix + block["text"])
            lines.append("")
        return ("\n".join(lines) + "\n").encode("utf-8")
    if fmt == "docx":
        from docx import Document
        document = Document()
        if title:
            document.add_heading(title, 0)
        for block in blocks:
            if block["type"] == "table":
                table = document.add_table(rows=1, cols=len(block["headers"]))
                table.style = "Light Shading Accent 1"
                for cell, value in zip(table.rows[0].cells, block["headers"]):
                    cell.text = str(value)
                for row in block["rows"]:
                    for cell, value in zip(table.add_row().cells, row):
                        cell.text = str(value)
            elif block["type"] == "heading":
                document.add_heading(block["text"], block.get("level", 2))
            else:
                document.add_paragraph(block["text"], "List Bullet" if block["type"] == "bullet" else None)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    if fmt == "pdf":
        import reportlab
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.colors import HexColor
        font = TTFont("CodaliVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"))
        pdfmetrics.registerFont(font)
        content = title + "".join(block.get("text", "") + "".join(map(str, block.get("headers", []))) + "".join(str(value) for row in block.get("rows", []) for value in row) for block in blocks)
        if any(ord(char) not in font.face.charToGlyph and char not in "\n\r\t" for char in content):
            raise ValueError("Bundled PDF font cannot represent some characters; choose DOCX/Markdown instead of losing text")
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "CodaliVera"
        def paragraph(value, style="BodyText"):
            return Paragraph(escape(str(value)).replace("\n", "<br/>"), styles[style])
        elements = [paragraph(title, "Title"), Spacer(1, 12)] if title else []
        for block in blocks:
            if block["type"] == "table":
                rows = [[paragraph(cell) for cell in row] for row in [block["headers"], *block["rows"]]]
                table = Table(rows, colWidths=[450 / len(block["headers"])] * len(block["headers"]), repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HexColor("#e8eef5")), ("GRID", (0, 0), (-1, -1), .5, HexColor("#9aaabb")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
                elements.append(table)
            else:
                value = ("• " if block["type"] == "bullet" else "") + block["text"]
                elements.append(paragraph(value, "Heading" + str(min(block.get("level", 2), 3)) if block["type"] == "heading" else "BodyText"))
            elements.append(Spacer(1, 8))
        output = io.BytesIO()
        SimpleDocTemplate(output, title=title, author="", invariant=1).build(elements)
        return output.getvalue()
    raise ValueError("Unsupported document format")


try:
    request = json.loads(Path("input.json").read_text())
    args = request["args"]
    fmt = args["format"]
    if request["operation"] == "read":
        data = Path("source").read_bytes()
        result = read_document(data, fmt, args)
    elif request["operation"] == "create":
        for block in args.get("blocks", []):
            if block["type"] == "table":
                if not block.get("headers") or any(len(row) != len(block["headers"]) for row in block.get("rows", [])):
                    raise ValueError("Every table row must match its nonempty header")
            elif "text" not in block:
                raise ValueError("Non-table blocks require text")
        data = create_document(fmt, args)
        if not data or len(data) > 8 * 1024 * 1024:
            raise ValueError("Generated document is empty or oversized")
        Path("output/artifact").write_bytes(data)
        result = {"format": fmt, "mimeType": MIME[fmt], "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), "verifiedStructure": read_document(data, fmt, {"maxChars": 1000, "maxRows": 5, "maxSheets": 2, "maxPages": 2})}
    else:
        raise ValueError("Unknown operation")
    print(json.dumps({"versions": VERSIONS, "resourceLimits": LIMITS, **result}, ensure_ascii=False, allow_nan=False))
except Exception as error:
    print(json.dumps({"error": str(error)[:1200], "versions": VERSIONS}))
